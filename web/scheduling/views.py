from datetime import timedelta
from io import BytesIO
from urllib.parse import quote

from django import forms as django_forms
from django.conf import settings
from django.contrib import messages
from django.contrib.auth import login
from django.contrib.auth.decorators import login_required, permission_required
from django.core.exceptions import PermissionDenied
from django.db.models import Count, Max, Prefetch, Q
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_POST
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .forms import (
    AssignmentForm,
    AssignmentResponseForm,
    MatchForm,
    RefereeSignupForm,
)
from .models import (
    Assignment,
    Competition,
    Match,
    RefereeProfile,
    RefereeReport,
)
from . import reports
from .report_docx import build_report_docx, report_filename
from .notifications import notify_assignment_removed, notify_leave_request
from .publishing import (
    after_match_update,
    match_snapshot,
    publish_match,
    publish_matches,
)


@never_cache
@login_required
def home(request):
    context = {
        "has_referee_profile": False,
        "pending_assignments": [],
        "upcoming_assignments": [],
        "unread_notifications": request.user.notifications.filter(
            read_at__isnull=True
        )[:5],
        "pending_reports": RefereeReport.objects.filter(
            author__user=request.user,
        ).exclude(
            status=RefereeReport.Status.SUBMITTED,
        ).select_related(
            "match__home_team", "match__away_team"
        ).order_by("match__kickoff_at"),
    }

    try:
        referee_profile = request.user.referee_profile
    except RefereeProfile.DoesNotExist:
        pass
    else:
        assignments = Assignment.objects.filter(
            referee=referee_profile,
            match__assignment_status=(
                Match.AssignmentStatus.PUBLISHED
            ),
            # 已取消、已结束或已经开球的比赛不再需要反馈。
            match__status=Match.Status.SCHEDULED,
            match__kickoff_at__gt=timezone.now(),
        ).select_related(
            "match",
            "match__competition",
            "match__home_team",
            "match__away_team",
            "match__venue",
        ).order_by(
            "match__kickoff_at",
            "position",
        )

        context.update(
            {
                "has_referee_profile": True,
                "pending_assignments": assignments.filter(
                    response_status=(
                        Assignment.ResponseStatus.PENDING
                    )
                ),
                "upcoming_assignments": assignments.exclude(
                    response_status=(
                        Assignment.ResponseStatus.PENDING
                    )
                ),
            }
        )

    if request.user.has_perm("scheduling.change_assignment"):
        context["leave_requests"] = Assignment.objects.filter(
            response_status=Assignment.ResponseStatus.LEAVE,
            match__kickoff_at__date__gte=timezone.localdate(),
        ).exclude(
            match__status=Match.Status.CANCELLED,
        ).select_related(
            "referee",
            "match",
            "match__home_team",
            "match__away_team",
            "match__venue",
        ).order_by(
            "match__kickoff_at",
            "position",
        )

    return render(
        request,
        "scheduling/home.html",
        context,
    )


@login_required
@permission_required("scheduling.view_match", raise_exception=True)
def match_list(request):
    today = timezone.localdate()
    can_see_drafts = request.user.has_perm("scheduling.change_assignment")

    matches = Match.objects.select_related(
        "competition",
        "home_team",
        "away_team",
        "venue",
    ).annotate(assigned_count=Count("assignments"))

    upcoming_matches = list(
        matches.filter(kickoff_at__date__gte=today)
        .prefetch_related(
            Prefetch(
                "assignments",
                queryset=Assignment.objects.select_related("referee"),
            )
        )
        .order_by("kickoff_at", "id")
    )

    # 每场比赛的四个岗位：草稿阶段只有排班管理员能看到名单（与比赛详情一致）。
    summary = {"matches": 0, "gaps": 0, "leave": 0, "pending": 0}
    for match in upcoming_matches:
        by_position = {
            assignment.position: assignment
            for assignment in match.assignments.all()
        }
        match.lineup_visible = (
            can_see_drafts
            or match.assignment_status == Match.AssignmentStatus.PUBLISHED
        )
        match.lineup = [
            {"label": label, "assignment": by_position.get(position)}
            for position, label in Assignment.Position.choices
        ]
        statuses = [a.response_status for a in by_position.values()]
        match.gap_count = (
            len(Assignment.Position.choices) - len(by_position)
            if match.status == Match.Status.SCHEDULED
            else 0
        )
        match.leave_count = statuses.count(Assignment.ResponseStatus.LEAVE)
        match.pending_count = (
            statuses.count(Assignment.ResponseStatus.PENDING)
            if match.assignment_status == Match.AssignmentStatus.PUBLISHED
            else 0
        )
        if match.status == Match.Status.SCHEDULED and match.lineup_visible:
            summary["matches"] += 1
            summary["gaps"] += match.gap_count
            summary["leave"] += match.leave_count
            summary["pending"] += match.pending_count

    filters = {
        "gaps": ("缺裁判", lambda m: m.gap_count),
        "leave": ("有请假", lambda m: m.leave_count),
        "pending": ("待确认", lambda m: m.pending_count),
    }
    active_filter = request.GET.get("show", "")
    if active_filter in filters:
        predicate = filters[active_filter][1]
        upcoming_matches = [
            match
            for match in upcoming_matches
            if match.lineup_visible
            and match.status == Match.Status.SCHEDULED
            and predicate(match)
        ]
    else:
        active_filter = ""

    context = {
        "upcoming_matches": upcoming_matches,
        "past_matches": matches.filter(
            kickoff_at__date__lt=today
        ).order_by("-kickoff_at", "-id"),
        "summary": summary,
        "show_summary": can_see_drafts,
        "filter_options": [("", "全部")]
        + [(key, label) for key, (label, _) in filters.items()],
        "active_filter": active_filter,
    }

    return render(request, "scheduling/match_list.html", context)


@login_required
@permission_required(
    "scheduling.publish_assignments",
    raise_exception=True,
)
def assignment_export(request):
    matches = Match.objects.filter(
        assignment_status=Match.AssignmentStatus.PUBLISHED
    ).select_related(
        "competition",
        "home_team",
        "away_team",
        "venue",
    ).prefetch_related(
        "assignments__referee",
    ).order_by(
        "kickoff_at",
        "id",
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "裁判安排"
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A4"

    headers = [
        "赛事",
        "赛季",
        "场次编号",
        "轮次",
        "开球时间",
        "主队",
        "客队",
        "比赛场地",
        "赛事级别",
        "比赛状态",
        "主裁判",
        "主裁反馈",
        "第一助理裁判",
        "第一助理反馈",
        "第二助理裁判",
        "第二助理反馈",
        "第四官员",
        "第四官员反馈",
        "安排发布时间",
    ]

    last_column = get_column_letter(len(headers))

    worksheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=len(headers),
    )
    title_cell = worksheet["A1"]
    title_cell.value = f"{settings.SITE_NAME}裁判安排表"
    title_cell.font = Font(
        name="微软雅黑",
        size=16,
        bold=True,
        color="FFFFFF",
    )
    title_cell.fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )
    title_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )
    worksheet.row_dimensions[1].height = 30

    worksheet.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=len(headers),
    )
    exported_at = timezone.localtime().strftime(
        "%Y年%m月%d日 %H:%M"
    )
    worksheet["A2"] = (
        f"仅包含已经发布的裁判安排　导出时间：{exported_at}"
    )
    worksheet["A2"].font = Font(
        name="微软雅黑",
        size=10,
        color="666666",
    )
    worksheet["A2"].alignment = Alignment(
        horizontal="left",
        vertical="center",
    )
    worksheet.row_dimensions[2].height = 22

    thin_side = Side(
        style="thin",
        color="D9E2F3",
    )
    thin_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    for column_number, header in enumerate(headers, start=1):
        cell = worksheet.cell(
            row=3,
            column=column_number,
            value=header,
        )
        cell.font = Font(
            name="微软雅黑",
            size=10,
            bold=True,
            color="FFFFFF",
        )
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="4472C4",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        cell.border = thin_border

    worksheet.row_dimensions[3].height = 26

    def local_excel_datetime(value):
        if value is None:
            return None

        if timezone.is_aware(value):
            value = timezone.localtime(value)

        return value.replace(tzinfo=None)

    def assignment_values(assignments, position):
        assignment = assignments.get(position)

        if assignment is None:
            return "", ""

        return (
            assignment.referee.name,
            assignment.get_response_status_display(),
        )

    for match in matches:
        assignments = {
            assignment.position: assignment
            for assignment in match.assignments.all()
        }

        referee_name, referee_status = assignment_values(
            assignments,
            Assignment.Position.REFEREE,
        )
        assistant_1_name, assistant_1_status = assignment_values(
            assignments,
            Assignment.Position.ASSISTANT_1,
        )
        assistant_2_name, assistant_2_status = assignment_values(
            assignments,
            Assignment.Position.ASSISTANT_2,
        )
        fourth_name, fourth_status = assignment_values(
            assignments,
            Assignment.Position.FOURTH_OFFICIAL,
        )

        worksheet.append(
            [
                match.competition.name,
                match.competition.season,
                match.match_number,
                match.round_name,
                local_excel_datetime(match.kickoff_at),
                match.home_team.name,
                match.away_team.name,
                match.venue.name,
                match.match_level,
                match.get_status_display(),
                referee_name,
                referee_status,
                assistant_1_name,
                assistant_1_status,
                assistant_2_name,
                assistant_2_status,
                fourth_name,
                fourth_status,
                local_excel_datetime(match.published_at),
            ]
        )

    for row in worksheet.iter_rows(
        min_row=4,
        max_row=worksheet.max_row,
        min_col=1,
        max_col=len(headers),
    ):
        for cell in row:
            cell.font = Font(
                name="微软雅黑",
                size=10,
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            cell.border = thin_border

        row[4].number_format = "yyyy-mm-dd hh:mm"
        row[18].number_format = "yyyy-mm-dd hh:mm"

    column_widths = [
        18,
        14,
        13,
        15,
        18,
        18,
        18,
        18,
        12,
        12,
        14,
        12,
        16,
        12,
        16,
        12,
        14,
        12,
        18,
    ]

    for column_number, width in enumerate(
        column_widths,
        start=1,
    ):
        worksheet.column_dimensions[
            get_column_letter(column_number)
        ].width = width

    worksheet.auto_filter.ref = (
        f"A3:{last_column}{worksheet.max_row}"
    )
    worksheet.print_title_rows = "1:3"
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True

    output = BytesIO()
    workbook.save(output)

    filename = (
        "裁判安排_"
        f"{timezone.localtime().strftime('%Y%m%d_%H%M')}.xlsx"
    )
    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = (
        "attachment; "
        f"filename*=UTF-8''{quote(filename)}"
    )
    response["Cache-Control"] = "no-store"

    return response


@login_required
@permission_required("scheduling.view_match", raise_exception=True)
def match_detail(request, pk):
    match = get_object_or_404(
        Match.objects.select_related(
            "competition",
            "home_team",
            "away_team",
            "venue",
            "created_by",
        ).prefetch_related(
            "assignments__referee",
        ),
        pk=pk,
    )

    # 草稿阶段只有排班管理员可以看到裁判名单。
    can_view_assignments = (
        match.assignment_status
        == Match.AssignmentStatus.PUBLISHED
        or request.user.has_perm("scheduling.change_assignment")
    )
    assignments_by_position = (
        {
            assignment.position: assignment
            for assignment in match.assignments.all()
        }
        if can_view_assignments
        else {}
    )
    assignment_rows = []

    for position, position_name in Assignment.Position.choices:
        assignment = assignments_by_position.get(position)

        assignment_rows.append(
            {
                "position_name": position_name,
                "assignment": assignment,
                "can_respond": (
                    assignment is not None
                    and match.response_closed_reason() is None
                    and assignment.referee.user_id
                    == request.user.id
                ),
            }
        )

    report = RefereeReport.objects.filter(match=match).select_related("author").first()
    is_manager = reports.is_report_manager(request.user)
    report_context = {
        "report": report,
        "can_edit_report": report is not None
        and reports.can_edit_report(request.user, report),
        "can_download_report": report is not None
        and reports.can_download_report(request.user, report),
        "can_start_report": report is None
        and reports.report_start_problem(request.user, match) is None,
        "report_request_form": (
            reports.ReportRequestForm(
                match=match,
                initial={
                    "author": match.assignments.order_by("position").first()
                },
            )
            if is_manager
            and match.assignment_status == Match.AssignmentStatus.PUBLISHED
            and match.status != Match.Status.CANCELLED
            and (report is None or not report.is_submitted)
            else None
        ),
        "show_report_section": report is not None or is_manager
        or reports.match_referee_profile(request.user, match) is not None,
    }

    return render(
        request,
        "scheduling/match_detail.html",
        {
            "match": match,
            "assignment_rows": assignment_rows,
            "can_view_assignments": can_view_assignments,
            **report_context,
        },
    )


@login_required
@permission_required("scheduling.add_match", raise_exception=True)
def match_create(request):
    form = MatchForm(
        request.POST if request.method == "POST" else None
    )

    if request.method == "POST" and form.is_valid():
        match = form.save(commit=False)
        match.created_by = request.user
        match.save()

        messages.success(request, "比赛已成功录入。")
        return redirect(
            "scheduling:match_detail",
            pk=match.pk,
        )

    return render(
        request,
        "scheduling/match_form.html",
        {
            "form": form,
            "page_title": "录入比赛",
            "submit_label": "保存比赛",
        },
    )


@login_required
@permission_required("scheduling.change_match", raise_exception=True)
def match_update(request, pk):
    match = get_object_or_404(Match, pk=pk)
    previous = match_snapshot(match)
    form = MatchForm(
        request.POST if request.method == "POST" else None,
        instance=match,
    )

    if request.method == "POST" and form.is_valid():
        match = form.save()
        notified = after_match_update(request, match, previous)
        if notified:
            messages.success(
                request,
                f"比赛信息已更新，已通知本场 {notified} 名裁判。",
            )
        else:
            messages.success(request, "比赛信息已更新。")

        return redirect(
            "scheduling:match_detail",
            pk=match.pk,
        )

    return render(
        request,
        "scheduling/match_form.html",
        {
            "form": form,
            "match": match,
            "page_title": "修改比赛",
            "submit_label": "保存修改",
        },
    )

@login_required
@permission_required(
    (
        "scheduling.add_assignment",
        "scheduling.change_assignment",
    ),
    raise_exception=True,
)
def assignment_update(request, pk):
    match = get_object_or_404(
        Match.objects.select_related(
            "competition",
            "home_team",
            "away_team",
            "venue",
        ),
        pk=pk,
    )
    form = AssignmentForm(
        request.POST if request.method == "POST" else None,
        match=match,
    )

    if request.method == "POST" and form.is_valid():
        was_published = (
            match.assignment_status
            == Match.AssignmentStatus.PUBLISHED
        )
        changed = form.save(assigned_by=request.user)

        if changed and was_published:
            notify_assignment_removed(
                request, match, form.removed_referees
            )
            messages.success(
                request,
                "裁判安排已保存，原安排已退回草稿，"
                "请检查后重新发布。",
            )
        elif changed:
            messages.success(request, "裁判安排已保存。")
        else:
            messages.info(request, "裁判安排没有变化。")

        return redirect(
            "scheduling:match_detail",
            pk=match.pk,
        )

    return render(
        request,
        "scheduling/assignment_form.html",
        {
            "form": form,
            "match": match,
        },
    )

@login_required
@permission_required(
    "scheduling.publish_assignments",
    raise_exception=True,
)
@require_POST
def assignment_publish(request, pk):
    match = get_object_or_404(Match, pk=pk)
    ok, notified, problem = publish_match(request, match)

    if ok:
        messages.success(
            request,
            f"裁判安排已经发布，已通知 {notified} 名裁判。",
        )
    elif match.assignment_status == Match.AssignmentStatus.PUBLISHED:
        messages.info(request, problem)
    else:
        messages.error(request, problem)

    return redirect(
        "scheduling:match_detail",
        pk=match.pk,
    )


@login_required
@permission_required(
    "scheduling.publish_assignments",
    raise_exception=True,
)
@require_POST
def assignment_bulk_publish(request):
    match_ids = [
        int(value)
        for value in request.POST.getlist("match_ids")
        if value.isdigit()
    ]

    if not match_ids:
        messages.error(request, "请先勾选要发布的比赛。")
        return redirect("scheduling:match_list")

    matches = Match.objects.filter(pk__in=match_ids).select_related(
        "home_team",
        "away_team",
    ).order_by("kickoff_at", "id")
    published, notified, skipped = publish_matches(request, matches)

    if published:
        messages.success(
            request,
            f"已发布 {published} 场比赛的裁判安排，"
            f"已通知 {notified} 名裁判。",
        )
    for reason in skipped:
        messages.error(request, f"未发布：{reason}")

    return redirect("scheduling:match_list")


@login_required
def assignment_respond(request, pk):
    assignment = get_object_or_404(
        Assignment.objects.select_related(
            "match",
            "match__competition",
            "match__home_team",
            "match__away_team",
            "match__venue",
            "referee",
            "referee__user",
        ),
        pk=pk,
    )

    if assignment.referee.user_id != request.user.id:
        raise PermissionDenied("不能反馈其他裁判的安排。")

    closed_reason = assignment.match.response_closed_reason()
    if closed_reason:
        messages.error(request, closed_reason)
        return redirect(
            "scheduling:match_detail",
            pk=assignment.match_id,
        )

    previous_status = assignment.response_status
    form = AssignmentResponseForm(
        request.POST if request.method == "POST" else None,
        instance=assignment,
    )

    if request.method == "POST" and form.is_valid():
        assignment = form.save(commit=False)
        assignment.responded_at = timezone.now()
        assignment.save(
            update_fields=[
                "response_status",
                "response_note",
                "responded_at",
                "updated_at",
            ]
        )

        if (
            assignment.response_status
            == Assignment.ResponseStatus.CONFIRMED
        ):
            messages.success(request, "已经确认参加本场执法。")
        else:
            if previous_status != Assignment.ResponseStatus.LEAVE:
                notify_leave_request(request, assignment)
            messages.success(
                request,
                "请假申请已经提交，已通知排班管理员。",
            )

        return redirect(
            "scheduling:match_detail",
            pk=assignment.match_id,
        )

    return render(
        request,
        "scheduling/assignment_response_form.html",
        {
            "assignment": assignment,
            "form": form,
        },
    )


def signup(request):
    if request.user.is_authenticated:
        return redirect("scheduling:home")

    form = RefereeSignupForm(
        request.POST if request.method == "POST" else None
    )

    if request.method == "POST" and form.is_valid():
        try:
            user = form.save()
        except django_forms.ValidationError as error:
            form.add_error("invite_code", error)
        else:
            login(
                request,
                user,
                backend="django.contrib.auth.backends.ModelBackend",
            )
            messages.success(request, "注册成功，欢迎加入！")
            return redirect("scheduling:home")

    return render(
        request,
        "registration/signup.html",
        {"form": form},
    )


@login_required
@permission_required(
    "scheduling.view_refereeprofile",
    raise_exception=True,
)
def referee_workload(request):
    competitions = Competition.objects.all()
    selected_competition = None
    competition_id = request.GET.get("competition", "")

    if competition_id.isdigit():
        selected_competition = competitions.filter(
            pk=competition_id
        ).first()

    # 草稿也计入，排班时需要考虑尚未发布的安排；已取消的比赛不计。
    counted = ~Q(assignments__match__status=Match.Status.CANCELLED)
    if selected_competition is not None:
        counted &= Q(
            assignments__match__competition=selected_competition
        )

    now = timezone.now()

    def count(extra=None):
        condition = counted if extra is None else counted & extra
        return Count("assignments", filter=condition)

    referees = (
        RefereeProfile.objects.annotate(
            total=count(),
            as_referee=count(
                Q(assignments__position=Assignment.Position.REFEREE)
            ),
            as_assistant=count(
                Q(
                    assignments__position__in=[
                        Assignment.Position.ASSISTANT_1,
                        Assignment.Position.ASSISTANT_2,
                    ]
                )
            ),
            as_fourth=count(
                Q(
                    assignments__position=(
                        Assignment.Position.FOURTH_OFFICIAL
                    )
                )
            ),
            confirmed=count(
                Q(
                    assignments__response_status=(
                        Assignment.ResponseStatus.CONFIRMED
                    )
                )
            ),
            on_leave=count(
                Q(
                    assignments__response_status=(
                        Assignment.ResponseStatus.LEAVE
                    )
                )
            ),
            upcoming=count(Q(assignments__match__kickoff_at__gte=now)),
            last_match_at=Max(
                "assignments__match__kickoff_at",
                filter=counted
                & Q(assignments__match__kickoff_at__lt=now),
            ),
        )
        .filter(Q(is_active=True) | Q(total__gt=0))
        .order_by("-total", "name")
    )

    return render(
        request,
        "scheduling/referee_workload.html",
        {
            "competitions": competitions,
            "selected_competition": selected_competition,
            "referees": referees,
        },
    )


@never_cache
@login_required
def notification_list(request):
    notifications = list(request.user.notifications.all()[:50])
    unread_ids = [
        notification.pk
        for notification in notifications
        if notification.read_at is None
    ]

    # 包括没显示出来的更早通知，否则页首未读数永远消不掉。
    request.user.notifications.filter(read_at__isnull=True).update(
        read_at=timezone.now()
    )

    return render(
        request,
        "scheduling/notification_list.html",
        {
            "notifications": notifications,
            "unread_ids": set(unread_ids),
            # 本页打开后全部已读，页首不再显示未读数。
            "unread_notification_count": 0,
        },
    )


def admin_login_redirect(request):
    """后台登录统一使用网站的登录页，保留 next 参数。"""
    query = request.GET.copy()
    query.setdefault("next", reverse("admin:index"))
    return redirect(f"{reverse('login')}?{query.urlencode()}")


# ---- 裁判报告 ----
def _report_match(pk):
    return get_object_or_404(
        Match.objects.select_related(
            "competition", "home_team", "away_team", "venue"
        ),
        pk=pk,
    )


@login_required
def report_edit(request, pk):
    match = _report_match(pk)
    report = RefereeReport.objects.filter(match=match).select_related("author__user").first()

    if report is None:
        # 本场裁判在赛后自行开始填写，成为填写人。
        problem = reports.report_start_problem(request.user, match)
        if problem:
            messages.error(request, problem)
            return redirect("scheduling:match_detail", pk=match.pk)
        report = reports.new_report(
            match, reports.match_referee_profile(request.user, match)
        )
    elif not reports.can_edit_report(request.user, report):
        raise PermissionDenied("只有填写裁判和排班管理员可以修改这份报告。")

    submitting = request.method == "POST" and "submit" in request.POST
    form = reports.RefereeReportForm(
        request.POST if request.method == "POST" else None,
        instance=report,
        submitting=submitting,
    )
    formsets = reports.build_row_formsets(
        report, request.POST if request.method == "POST" else None
    )

    if (
        request.method == "POST"
        and form.is_valid()
        and all(
            formset.is_valid()
            for formset in reports.formsets_to_validate(formsets, request.POST)
        )
    ):
        report = form.save(commit=False)
        reports.apply_row_formsets(report, formsets)
        if submitting:
            report.status = RefereeReport.Status.SUBMITTED
            report.submitted_at = timezone.now()
        elif report.status != RefereeReport.Status.SUBMITTED:
            report.status = RefereeReport.Status.DRAFT
        report.save()

        if submitting:
            reports.notify_report_submitted(request, report)
            messages.success(request, "裁判报告已提交，已通知排班管理员。")
            return redirect("scheduling:match_detail", pk=match.pk)
        messages.success(request, "已保存。")
        return redirect("scheduling:report_edit", pk=match.pk)

    return render(
        request,
        "scheduling/report_form.html",
        {
            "match": match,
            "report": report,
            "form": form,
            "row_sections": [
                (key, label, formsets[key])
                for key, label, _formset in reports.ROW_SECTIONS
            ],
            "shootout_section": (
                reports.SHOOTOUT_SECTION[0],
                reports.SHOOTOUT_SECTION[1],
                formsets[reports.SHOOTOUT_SECTION[0]],
            ),
        },
    )


@login_required
@require_POST
def report_request(request, pk):
    if not reports.is_report_manager(request.user):
        raise PermissionDenied("只有排班管理员可以发起裁判报告要求。")
    match = _report_match(pk)
    report = RefereeReport.objects.filter(match=match).first()

    if report is not None and report.is_submitted:
        messages.info(request, "本场裁判报告已经提交。")
        return redirect("scheduling:match_detail", pk=match.pk)
    if match.assignment_status != Match.AssignmentStatus.PUBLISHED:
        messages.error(request, "请先发布本场裁判安排。")
        return redirect("scheduling:match_detail", pk=match.pk)

    form = reports.ReportRequestForm(request.POST, match=match)
    if not form.is_valid():
        messages.error(request, "请选择填写裁判。")
        return redirect("scheduling:match_detail", pk=match.pk)

    author = form.cleaned_data["author"].referee
    if report is None:
        report = reports.new_report(match, author, requested_by=request.user)
    else:
        report.author = author
        report.status = RefereeReport.Status.REQUESTED
        report.requested_by = request.user
        report.requested_at = timezone.now()
    report.save()
    reports.notify_report_requested(request, report)
    messages.success(request, f"已通知 {author.name} 提交裁判报告。")
    return redirect("scheduling:match_detail", pk=match.pk)


@login_required
def report_download(request, pk):
    report = get_object_or_404(
        RefereeReport.objects.select_related(
            "author",
            "match__competition",
            "match__home_team",
            "match__away_team",
            "match__venue",
        ),
        match_id=pk,
    )
    if not reports.can_download_report(request.user, report):
        raise PermissionDenied("没有权限下载这份裁判报告。")

    response = HttpResponse(
        build_report_docx(report),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
    )
    response["Content-Disposition"] = (
        f"attachment; filename*=UTF-8''{quote(report_filename(report))}"
    )
    response["Cache-Control"] = "no-store"
    return response


@login_required
def report_list(request):
    """排班管理员看全部报告；裁判看自己负责的报告和可以补交的比赛。"""
    manager = reports.is_report_manager(request.user)
    queryset = RefereeReport.objects.select_related(
        "author", "match__home_team", "match__away_team", "match__competition"
    ).order_by("-match__kickoff_at")
    if not manager:
        queryset = queryset.filter(author__user=request.user)

    status = request.GET.get("status", "")
    if status in RefereeReport.Status.values:
        queryset = queryset.filter(status=status)

    startable = []
    if not manager or hasattr(request.user, "referee_profile"):
        # 最近两周本场执法、已开球、还没有报告的比赛，可以直接补交。
        startable = (
            Match.objects.filter(
                assignments__referee__user=request.user,
                assignment_status=Match.AssignmentStatus.PUBLISHED,
                kickoff_at__lte=timezone.now(),
                kickoff_at__gte=timezone.now() - timedelta(days=14),
                report__isnull=True,
            )
            .exclude(status=Match.Status.CANCELLED)
            .select_related("home_team", "away_team")
            .order_by("-kickoff_at")
        )

    return render(
        request,
        "scheduling/report_list.html",
        {
            "reports": queryset,
            "startable": startable,
            "manager": manager,
            "status_options": [("", "全部")] + list(RefereeReport.Status.choices),
            "active_status": status if status in RefereeReport.Status.values else "",
        },
    )
