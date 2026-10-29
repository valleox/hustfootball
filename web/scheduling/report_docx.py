"""把裁判报告导出为 Word 文件，版式与协会的「华中科技大学裁判报告表」一致。"""

from io import BytesIO
from pathlib import Path

from django.utils import timezone
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .models import Assignment

TITLE = "华中科技大学裁判报告表"
BODY_SIZE = Pt(12)
HEADING_SIZE = Pt(14)
TABLE_SIZE = Pt(10.5)
ROW_HEIGHT = Cm(0.9)
TEXT_WIDTH = Cm(21 - 3.175 * 2)
# 警告、罚令出场表的列宽：队名、号码、姓名、时间、原因
DISCIPLINE_WIDTHS = [Cm(2.3), Cm(1.5), Cm(2.2), Cm(1.6), Cm(7.05)]
SHOOTOUT_WIDTHS = [Cm(1.6), Cm(3.6), Cm(2.4), Cm(4.0), Cm(3.05)]
ASSETS = Path(__file__).resolve().parent / "report_assets"


def _set_fonts(run, size):
    run.font.size = size
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")


def _add_run(paragraph, text, size=BODY_SIZE, bold=False, underline=False):
    run = paragraph.add_run(text)
    _set_fonts(run, size)
    run.bold = bold
    run.underline = underline
    return run


def _line(document, parts):
    """parts 为 (文字, 是否为填写内容) 列表；填写内容加下划线，与样表一致。"""
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for text, is_value in parts:
        _add_run(paragraph, text, underline=is_value)
    return paragraph


def _heading(document, text):
    paragraph = document.add_paragraph()
    _add_run(paragraph, text, size=HEADING_SIZE, bold=True)
    return paragraph


def _cell_text(cell, text, bold=False):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(paragraph, str(text or ""), size=TABLE_SIZE, bold=bold)


def _set_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width


def _grid_table(document, headers, rows, widths, min_rows=1):
    rows = list(rows) or []
    while len(rows) < min_rows:
        rows.append([""] * len(headers))
    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for column, header in enumerate(headers):
        _cell_text(table.cell(0, column), header)
    for row_index, values in enumerate(rows, start=1):
        table.rows[row_index].height = ROW_HEIGHT
        for column, value in enumerate(values):
            _cell_text(table.cell(row_index, column), value)
    _set_widths(table, widths)
    return table


def _notes_box(document, text):
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.text = ""
    lines = (text or "").splitlines() or [""]
    for index, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        # 与样表一致：每段首行缩进两个字。
        paragraph.paragraph_format.first_line_indent = TABLE_SIZE * 2
        paragraph.paragraph_format.line_spacing = 1.3
        _add_run(paragraph, line.strip(), size=TABLE_SIZE)
    # 内容很少时保留一点书写空间。
    table.rows[0].height = Cm(1.6)
    _set_widths(table, [TEXT_WIDTH])
    return table


def _title(document):
    """标题行：左侧校徽、中间标题、右侧会徽（无边框表格排版）。"""
    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, middle, right = table.rows[0].cells
    for cell, image, width in (
        (left, "hust_seal.png", Cm(2.4)),
        (right, "hustfa_badge.png", Cm(1.9)),
    ):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(ASSETS / image), width=width)
    middle.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = middle.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(paragraph, TITLE, size=Pt(22))
    _set_widths(table, [Cm(2.7), TEXT_WIDTH - Cm(5.4), Cm(2.7)])


def _minute(value):
    value = str(value or "").strip()
    return f"{value}min" if value.isdigit() else value


def team_name(team):
    return team.short_name or team.name


def officials_line(match):
    names = {
        assignment.position: assignment.referee.name
        for assignment in match.assignments.select_related("referee")
    }
    labels = (
        (Assignment.Position.REFEREE, "裁判员"),
        (Assignment.Position.ASSISTANT_1, "第一助理"),
        (Assignment.Position.ASSISTANT_2, "第二助理"),
        (Assignment.Position.FOURTH_OFFICIAL, "第四官员"),
    )
    return " ".join(f"{label}：{names.get(position, '')}" for position, label in labels)


def build_report_docx(report):
    match = report.match
    kickoff = timezone.localtime(match.kickoff_at)
    home, away = team_name(match.home_team), team_name(match.away_team)

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(3.175)
    section.top_margin = section.bottom_margin = Cm(2.54)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")

    _title(document)
    document.add_paragraph()

    score = ""
    if report.home_score is not None and report.away_score is not None:
        score = f"{report.home_score}:{report.away_score}"

    _line(document, [("1，比赛名称：", False), (report.competition_title or match.competition.name, True)])
    _line(
        document,
        [
            ("2，比赛对阵双方：", False),
            (home, True),
            ("VS", False),
            (away, True),
            ("  场序：", False),
            (match.match_number or "", True),
        ],
    )
    _line(
        document,
        [
            ("3，比赛时间：", False),
            (str(kickoff.year), False),
            ("年", False),
            (str(kickoff.month), True),
            ("月", False),
            (str(kickoff.day), True),
            ("日", False),
            (f"{kickoff.hour:02d}", True),
            ("时", False),
            (f"{kickoff.minute:02d}", True),
            ("分", False),
        ],
    )
    _line(document, [("4，比赛地点：", False), (match.venue.name, True)])
    result = [("5，比赛结果：80分钟内比分：", False), (score, True)]
    if report.has_shootout:
        shootout_score = ""
        if report.home_shootout_score is not None and report.away_shootout_score is not None:
            shootout_score = f"{report.home_shootout_score}:{report.away_shootout_score}"
        result += [("    点球决胜：", False), (shootout_score, True)]
    result += [("      获胜队：", False), (report.winner, True)]
    _line(document, result)

    _heading(document, "警告队员")
    _grid_table(
        document,
        ["队名", "号码", "队员姓名", "时间", "原因"],
        (
            [row.get("team"), row.get("number"), row.get("name"), _minute(row.get("minute")), row.get("reason")]
            for row in report.cautions
        ),
        DISCIPLINE_WIDTHS,
    )

    _heading(document, "罚令出场")
    _grid_table(
        document,
        ["队名", "号码", "队员姓名", "时间", "原因（必须写清楚）"],
        (
            [row.get("team"), row.get("number"), row.get("name"), _minute(row.get("minute")), row.get("reason")]
            for row in report.sendoffs
        ),
        DISCIPLINE_WIDTHS,
    )

    _heading(document, "进球队员")
    goal_rows = max(len(report.home_goals), len(report.away_goals), 1)
    goals = document.add_table(rows=2 + goal_rows, cols=6)
    goals.style = "Table Grid"
    goals.alignment = WD_TABLE_ALIGNMENT.CENTER
    goals.cell(0, 0).merge(goals.cell(0, 2))
    goals.cell(0, 3).merge(goals.cell(0, 5))
    _cell_text(goals.cell(0, 0), "主    队")
    _cell_text(goals.cell(0, 3), "客    队")
    for column, header in enumerate(["姓名", "号码", "时间"] * 2):
        _cell_text(goals.cell(1, column), header)
    for row in goals.rows:
        row.height = ROW_HEIGHT
    _set_widths(goals, [TEXT_WIDTH / 6] * 6)
    for offset, side in ((0, report.home_goals), (3, report.away_goals)):
        for index in range(goal_rows):
            row = side[index] if index < len(side) else {}
            _cell_text(goals.cell(2 + index, offset), row.get("name"))
            _cell_text(goals.cell(2 + index, offset + 1), row.get("number"))
            _cell_text(goals.cell(2 + index, offset + 2), _minute(row.get("minute")))

    if report.has_shootout:
        _heading(document, "点球决胜")
        _grid_table(
            document,
            ["顺序", "队名", "号码", "队员姓名", "结果"],
            (
                [index, row.get("team"), row.get("number"), row.get("name"), row.get("result")]
                for index, row in enumerate(report.shootout_kicks, start=1)
            ),
            SHOOTOUT_WIDTHS,
        )

    _heading(document, "红牌情况说明（如本页不够请另附页说明）")
    _notes_box(document, report.sendoff_notes)
    _heading(document, "罚球点球情况说明")
    _notes_box(document, report.penalty_notes)
    _heading(document, "严重错漏判情况说明")
    _notes_box(document, report.error_notes)
    _heading(document, "整场吹罚问题总结")
    _notes_box(document, report.summary)

    document.add_paragraph()
    _line(document, [(officials_line(match), False)])

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def report_filename(report):
    """例如：26华工杯 A01 软件vs公管 吴扬.docx（与协会的命名习惯接近）。"""
    match = report.match
    parts = [
        match.competition.name,
        match.match_number,
        f"{team_name(match.home_team)}vs{team_name(match.away_team)}",
        report.author.name,
    ]
    name = " ".join(part for part in parts if part)
    for char in '\\/:*?"<>|':
        name = name.replace(char, "_")
    return f"{name}.docx"
